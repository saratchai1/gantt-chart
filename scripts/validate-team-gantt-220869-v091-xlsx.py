#!/usr/bin/env python3
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
XLSX = DATA / "huai-kha-khaeng-team-gantt-220869-v0.9.1-thai.xlsx"
REPORT = DATA / "team-gantt-xlsx-validation-220869-v091.json"

EXPECTED_SHEETS = ["Gantt Chart", "Activity Data", "TBC Register", "Revision & Notes"]
EXPECTED_ACTIVITY_COUNT = 107
EXPECTED_TBC = 6
EXPECTED_CONTROL = 11
EXPECTED_CRITICAL = 10
EXPECTED_MONTHS = 40


def main() -> None:
    errors = []
    advisories = []
    if not XLSX.exists():
        raise SystemExit(f"Missing workbook: {XLSX}")

    wb = load_workbook(XLSX, data_only=False)
    if wb.sheetnames != EXPECTED_SHEETS:
        errors.append({"issue": "sheet_names", "expected": EXPECTED_SHEETS, "actual": wb.sheetnames})

    gantt = wb["Gantt Chart"]
    data = wb["Activity Data"]
    tbc = wb["TBC Register"]
    rev = wb["Revision & Notes"]

    # Flat data checks.
    activity_ids = []
    timing_statuses = []
    match_levels = []
    critical_values = []
    formula_count = 0
    for row in range(2, data.max_row + 1):
        aid = data.cell(row=row, column=2).value
        if aid:
            activity_ids.append(str(aid))
        timing_statuses.append(str(data.cell(row=row, column=14).value or ""))
        match_levels.append(str(data.cell(row=row, column=12).value or ""))
        critical_values.append(str(data.cell(row=row, column=25).value or ""))
        elapsed = data.cell(row=row, column=17).value
        if isinstance(elapsed, str) and elapsed.startswith("="):
            formula_count += 1

    if len(activity_ids) != EXPECTED_ACTIVITY_COUNT:
        errors.append({"issue": "activity_count", "expected": EXPECTED_ACTIVITY_COUNT, "actual": len(activity_ids)})
    if len(set(activity_ids)) != len(activity_ids):
        duplicates = sorted({aid for aid in activity_ids if activity_ids.count(aid) > 1})
        errors.append({"issue": "duplicate_activity_ids", "sample": duplicates[:20]})
    if formula_count != EXPECTED_ACTIVITY_COUNT:
        errors.append({"issue": "elapsed_span_formulas", "expected": EXPECTED_ACTIVITY_COUNT, "actual": formula_count})

    tbc_count = sum(value == "TBC_TEAM_CONFIRMATION" for value in timing_statuses)
    control_count = sum(value == "CONTROL_STREAM_MATCH" for value in match_levels)
    critical_count = sum(value == "CONTAINS_ZERO_FLOAT_DETAIL" for value in critical_values)
    if tbc_count != EXPECTED_TBC:
        errors.append({"issue": "tbc_count", "expected": EXPECTED_TBC, "actual": tbc_count})
    if control_count != EXPECTED_CONTROL:
        errors.append({"issue": "control_window_count", "expected": EXPECTED_CONTROL, "actual": control_count})
    if critical_count != EXPECTED_CRITICAL:
        errors.append({"issue": "critical_exposure_count", "expected": EXPECTED_CRITICAL, "actual": critical_count})

    # Gantt timeline checks.
    expected_last_col = 12 + EXPECTED_MONTHS
    if gantt.max_column != expected_last_col:
        errors.append({"issue": "timeline_columns", "expected": expected_last_col, "actual": gantt.max_column})
    if gantt.freeze_panes != "M6":
        errors.append({"issue": "freeze_panes", "expected": "M6", "actual": str(gantt.freeze_panes)})
    if gantt["A1"].value != "แผนงานก่อสร้างห้วยขาแข้ง — Gantt Chart ตามกิจกรรมทีมงาน 220869":
        errors.append({"issue": "title"})

    # Count activity IDs in the Gantt sheet exactly once.
    gantt_ids = []
    for row in range(6, gantt.max_row + 1):
        value = gantt.cell(row=row, column=2).value
        if isinstance(value, str) and value.startswith("T220869-"):
            gantt_ids.append(value)
    if len(gantt_ids) != EXPECTED_ACTIVITY_COUNT:
        errors.append({"issue": "gantt_activity_count", "expected": EXPECTED_ACTIVITY_COUNT, "actual": len(gantt_ids)})
    if sorted(gantt_ids) != sorted(activity_ids):
        missing = sorted(set(activity_ids) - set(gantt_ids))
        extra = sorted(set(gantt_ids) - set(activity_ids))
        errors.append({"issue": "gantt_activity_id_mismatch", "missing": missing[:20], "extra": extra[:20]})

    # TBC handoff sheet checks.
    tbc_rows = max(0, tbc.max_row - 1)
    if tbc_rows != EXPECTED_TBC:
        errors.append({"issue": "tbc_register_rows", "expected": EXPECTED_TBC, "actual": tbc_rows})

    # Formula-driven revision summary should remain formulas.
    summary_formulas = 0
    for row in range(1, rev.max_row + 1):
        value = rev.cell(row=row, column=2).value
        if isinstance(value, str) and value.startswith("="):
            summary_formulas += 1
    if summary_formulas < 5:
        errors.append({"issue": "revision_summary_formulas", "expected_minimum": 5, "actual": summary_formulas})

    # Basic file health / readability.
    size = XLSX.stat().st_size
    if size < 25000:
        errors.append({"issue": "file_size_too_small", "bytes": size})
    if size > 10_000_000:
        advisories.append({"issue": "file_size_large", "bytes": size})

    status = "PASS" if not errors and not advisories else "PASS_WITH_ADVISORIES" if not errors else "FAIL"
    report = {
        "status": status,
        "workbook": str(XLSX.relative_to(ROOT)),
        "size_bytes": size,
        "sheets": wb.sheetnames,
        "activity_rows": len(activity_ids),
        "gantt_activity_rows": len(gantt_ids),
        "tbc_rows": tbc_count,
        "control_window_rows": control_count,
        "critical_exposure_rows": critical_count,
        "elapsed_span_formula_rows": formula_count,
        "timeline_months": EXPECTED_MONTHS,
        "errors": errors,
        "advisories": advisories,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
